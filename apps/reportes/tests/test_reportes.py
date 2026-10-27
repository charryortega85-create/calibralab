import csv
import io
from datetime import date, timedelta

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.plan.models import ItemPlan, Lote, PlanAnual


def test_hoja_de_vida_en_pdf(cliente_de, consulta, crear_equipo, calibrar):
    e = crear_equipo()
    calibrar(e, date(2026, 3, 9))
    c = cliente_de(consulta)
    r = c.get(reverse("reportes:hoja_vida", args=[e.pk]))
    assert r["Content-Type"] == "application/pdf" and r.content[:4] == b"%PDF"
    html = c.get(reverse("reportes:hoja_vida", args=[e.pk]), {"formato": "html"}).content.decode()
    assert "CC-1" in html and "Cambios de estado" in html and "Mantenimiento" not in html


def test_hoja_de_vida_sin_biblioteca_de_pdf_se_abre_para_imprimir(cliente_de, consulta, crear_equipo, monkeypatch):
    import sys

    monkeypatch.setitem(sys.modules, "weasyprint", None)  # como en Windows sin GTK
    e = crear_equipo()
    r = cliente_de(consulta).get(reverse("reportes:hoja_vida", args=[e.pk]))
    html = r.content.decode()
    assert r.status_code == 200 and "Guardar como PDF" in html and "window.print()" in html and "EQ-0001" in html


def test_estado_exporta_excel(cliente_de, consulta, crear_equipo, calibrar):
    calibrar(crear_equipo("EQ-0001"), date(2020, 1, 1))
    crear_equipo("EQ-0002")
    r = cliente_de(consulta).get(reverse("reportes:estado"), {"formato": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    filas = list(ws.values)
    assert filas[0][0] == "Código" and [f[0] for f in filas[1:]] == ["EQ-0001", "EQ-0002"]  # vencido antes que sin registro
    assert filas[1][6] == "Vencido"


def test_vencimientos_por_tramo(cliente_de, consulta, crear_equipo, calibrar):
    hoy = timezone.localdate()
    for codigo, dias in [("EQ-10", 10), ("EQ-45", 45), ("EQ-80", 80), ("EQ-200", 200)]:
        calibrar(crear_equipo(codigo, intervalo_meses=1), hoy + timedelta(days=dias) - timedelta(days=31))
    c = cliente_de(consulta)

    def codigos(dias):
        r = c.get(reverse("reportes:vencimientos"), {"dias": dias, "formato": "csv", "vencidos": "0"})
        return [f[0] for f in csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";")][1:]

    assert len(codigos(30)) == 1 and len(codigos(60)) == 2 and len(codigos(90)) == 3


def test_ejecucion_y_presupuesto(cliente_de, consulta):
    plan = PlanAnual.objects.create(ejercicio=2026)
    lote = Lote.objects.create(codigo="L-01", nombre="Balanzas", trabajos="x")
    ItemPlan.objects.create(plan=plan, lote=lote, mes_previsto=3, monto_previsto=1000, monto_pagado=500,
                            estado_ejecucion="ejecutado")
    c = cliente_de(consulta)
    r = c.get(reverse("reportes:ejecucion"), {"plan": plan.pk, "formato": "xlsx"})
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Ejecución por lote", "Avance presupuestario"]
    assert list(wb["Avance presupuestario"].values)[1] == ("Marzo", 1, 1, 1000, 500, 1000, 500, 50)
    assert "L-01" in c.get(reverse("reportes:ejecucion")).content.decode()
    r = c.get(reverse("reportes:ejecucion"), {"plan": plan.pk, "formato": "csv", "hoja": "presupuesto"})
    assert "Marzo" in r.content.decode("utf-8-sig")


def test_indice_de_reportes(cliente_de, consulta):
    assert "Vencimientos" in cliente_de(consulta).get(reverse("reportes:indice")).content.decode()
