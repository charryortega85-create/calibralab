"""Exportación a CSV y Excel (RF-12, RF-37, RF-38)."""
import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def _celda(valor):
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    if isinstance(valor, (int, float, Decimal, date, datetime)):
        return valor
    return str(valor)


def respuesta_csv(nombre, encabezados, filas):
    r = HttpResponse(content_type="text/csv; charset=utf-8")
    r["Content-Disposition"] = f'attachment; filename="{nombre}.csv"'
    r.write("﻿")  # BOM: Excel abre bien los acentos
    w = csv.writer(r, delimiter=";")
    w.writerow(encabezados)
    for fila in filas:
        w.writerow(["" if v is None else ("Sí" if v is True else "No" if v is False else v) for v in fila])
    return r


def libro_excel(hojas):
    """hojas: lista de (titulo, encabezados, filas)."""
    wb = Workbook()
    wb.remove(wb.active)
    for titulo, encabezados, filas in hojas:
        ws = wb.create_sheet(titulo[:31])
        ws.append(list(encabezados))
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0F766E")
        for fila in filas:
            ws.append([_celda(v) for v in fila])
        for i, enc in enumerate(encabezados, start=1):
            largo = max([len(str(enc))] + [len(str(_celda(f[i - 1]))) for f in filas[:200]] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 50)
        ws.freeze_panes = "A2"
    return wb


def respuesta_excel(nombre, encabezados, filas, hojas_extra=()):
    filas = list(filas)
    wb = libro_excel([(nombre, encabezados, filas), *hojas_extra])
    buf = BytesIO()
    wb.save(buf)
    r = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    r["Content-Disposition"] = f'attachment; filename="{nombre}.xlsx"'
    return r


def exportar(formato, nombre, encabezados, filas):
    filas = list(filas)
    if formato == "csv":
        return respuesta_csv(nombre, encabezados, filas)
    return respuesta_excel(nombre, encabezados, filas)
