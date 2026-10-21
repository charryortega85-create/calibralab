from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.catalogos.models import TipoInstrumento
from apps.equipos import importar as imp
from apps.equipos.models import Equipo


def libro(hojas):
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        ws.append(imp.encabezados_de(nombre))
        for f in filas:
            ws.append([f.get(c, "") for c in imp.encabezados_de(nombre)])
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("e.xlsx", buf.getvalue())


def fila(codigo, tipo, **extra):
    d = {"codigo": codigo, "descripcion": "Equipo", "tipo": tipo, "marca": "Marca A", "modelo": "M1",
         "numero_serie": f"S-{codigo}", "ubicacion": "Sala de balanzas", "responsable": "tecnico", "unidad": "g",
         "criticidad": "Alta", "intervalo_meses": 12}
    d.update(extra)
    return d


def test_plantilla_trae_una_hoja_por_familia(cliente_de, tecnico):
    r = cliente_de(tecnico).get(reverse("equipos:plantilla"))
    wb = load_workbook(BytesIO(r.content))
    assert set(imp.HOJAS) <= set(wb.sheetnames)
    assert [c.value for c in wb["balanzas"][1]][-3:] == ["clase_exactitud", "division_d", "division_e"]
    assert "Balanza analítica" in [c.value for c in wb["listas"]["B"]]


def test_hojas_de_familia_guardan_los_datos_tecnicos(cliente_de, tecnico, catalogo):
    c = cliente_de(tecnico)
    archivo = libro({
        "balanzas": [fila("EQ-0700", "Balanza analítica", clase_exactitud="i", division_d="0,0001", critico="Sí")],
        "pipetas": [fila("EQ-0701", "Pipeta multicanal", canales="8", tipo_volumen="variable")],
        "equipos": [fila("EQ-0702", "")],
    })
    r = c.post(reverse("equipos:importar"), {"archivo": archivo})
    assert "Guardar 3 equipos" in r.content.decode()
    c.post(reverse("equipos:importar"), {"confirmar": "1"})
    b = Equipo.objects.get(codigo="EQ-0700")
    assert b.tipo.nombre == "Balanza analítica" and b.critico
    assert b.caracteristicas == {"clase_exactitud": "I", "division_d": "0.0001"}
    assert Equipo.objects.get(codigo="EQ-0701").caracteristicas["tipo_volumen"] == "Variable"
    assert Equipo.objects.get(codigo="EQ-0702").tipo is None


def test_tipo_de_otra_familia_o_vacio_es_error(cliente_de, tecnico, catalogo):
    archivo = libro({"balanzas": [fila("EQ-0710", "Estufa de cultivo"), fila("EQ-0711", "")],
                     "pipetas": [fila("EQ-0712", "Pipeta monocanal", canales="5")]})
    r = cliente_de(tecnico).post(reverse("equipos:importar"), {"archivo": archivo})
    html = r.content.decode()
    assert "balanzas, fila 2" in html and "balanzas, fila 3" in html and "pipetas, fila 2" in html
    assert "Canales: elige 1, 8, 12, 16." in html
    assert not Equipo.objects.exists()


def test_administrador_agrega_un_tipo_nuevo_a_la_familia_de_la_hoja(cliente_de, admin, tecnico, catalogo):
    c = cliente_de(admin)
    archivo = libro({"termicos": [fila("EQ-0720", "Baño termostático", setpoints="37; 56", unidad="g")]})
    c.post(reverse("equipos:importar"), {"archivo": archivo, "crear_faltantes": "on"})
    c.post(reverse("equipos:importar"), {"confirmar": "1"})
    t = TipoInstrumento.objects.get(nombre="Baño termostático")
    assert t.familia == "termico" and Equipo.objects.get(codigo="EQ-0720").tipo == t
