"""
RF-12: carga de varios equipos a la vez desde Excel o CSV.

Primero se revisan todas las filas con las mismas reglas del formulario de alta y se
muestran los errores; solo si no hay ninguno se guardan todas juntas (RNF-10).

El Excel trae una hoja por familia de instrumentos (balanzas, térmicos, cabinas, lectores,
pipetas) con sus datos técnicos propios, y la hoja «equipos» para el resto.
"""
import csv
import io
from dataclasses import dataclass, field

from openpyxl import load_workbook

from apps.catalogos.models import Area, Familia, Laboratorio, Marca, TipoInstrumento, Ubicacion, Unidad
from apps.cuentas.models import Usuario

from . import familias
from .forms import EquipoForm
from .models import Criticidad

# (encabezado en el archivo, campo del formulario, descripción para la plantilla)
COLUMNAS = [
    ("codigo", "codigo", "Código interno único, por ejemplo EQ-0001"),
    ("patrimonio", "patrimonio", "N.º de patrimonio (opcional)"),
    ("descripcion", "descripcion", "Descripción del equipo"),
    ("tipo", "tipo", "Tipo de instrumento, por ejemplo Balanza analítica (en cada hoja se elige de una lista)"),
    ("marca", "marca", "Nombre de la marca, tal como figura en la lista de marcas"),
    ("modelo", "modelo", "Modelo"),
    ("numero_serie", "numero_serie", "N.º de serie"),
    ("laboratorio", "laboratorio", "Nombre del laboratorio, tal como figura en la lista"),
    ("area", "area", "Nombre del área dentro de ese laboratorio"),
    ("ubicacion", "ubicacion", "Nombre de la ubicación, tal como figura en la lista"),
    ("responsable", "responsable", "Usuario del responsable, por ejemplo mperez"),
    ("unidad", "unidad", "Símbolo de la unidad de medida, por ejemplo g o °C"),
    ("rango_min", "rango_min", "Rango desde (opcional)"),
    ("rango_max", "rango_max", "Rango hasta (opcional)"),
    ("resolucion", "resolucion", "Resolución (opcional)"),
    ("capacidad", "capacidad", "Capacidad o alcance (opcional)"),
    ("tolerancia", "tolerancia", "Tolerancia (opcional)"),
    ("criticidad", "criticidad", "Alta, Media o Baja"),
    ("critico", "critico", "Sí si es un equipo crítico (afecta la validez de los resultados), si no No"),
    ("intervalo_meses", "intervalo_meses", "Intervalo de calibración en meses"),
]
ENCABEZADOS = [c[0] for c in COLUMNAS]

# Hoja del Excel → familia. En la hoja «equipos» va cualquier tipo, o ninguno.
HOJAS = {"equipos": None, "balanzas": Familia.BALANZA, "termicos": Familia.TERMICO, "cabinas": Familia.CABINA,
         "lectores": Familia.LECTOR, "pipetas": Familia.PIPETA}
SI = ("si", "sí", "x", "1", "true", "verdadero")


def encabezados_de(hoja):
    familia = HOJAS.get(hoja)
    return ENCABEZADOS + [d.clave for d in familias.datos(familia)]


def _hoja(nombre):
    n = (nombre or "").strip().lower().replace("é", "e").replace("á", "a")
    return n if n in HOJAS else None


@dataclass
class Fila:
    numero: int
    datos: dict
    errores: list = field(default_factory=list)
    avisos: list = field(default_factory=list)   # lo que se agregará a las listas al guardar

    @property
    def codigo(self):
        return self.datos.get("codigo", "")

    @property
    def lugar(self):
        hoja, n = self.numero if isinstance(self.numero, (list, tuple)) else ("", self.numero)
        return f"{hoja}, fila {n}" if hoja else f"fila {n}"


def _registros(filas, hoja):
    if not filas:
        return []
    enc = [str(c or "").strip().lower() for c in filas[0]]
    salida = []
    for n, valores in enumerate(filas[1:], start=2):
        if not any(v not in (None, "") for v in valores):
            continue
        reg = {enc[i]: ("" if v is None else _texto(v)) for i, v in enumerate(valores) if i < len(enc)}
        reg["_hoja"], reg["_fila"] = hoja, n
        salida.append(reg)
    return salida


def _texto(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def leer(archivo):
    """
    Devuelve una lista de diccionarios con los encabezados en minúsculas, más la hoja y la fila.
    Un Excel con hojas de la plantilla se lee hoja por hoja; uno sin ellas (o un CSV), como «equipos».
    """
    nombre = archivo.name.lower()
    if nombre.endswith(".xlsx"):
        wb = load_workbook(archivo, read_only=True, data_only=True)
        hojas = [(h, _hoja(ws.title)) for h, ws in ((ws, ws) for ws in wb.worksheets)]
        conocidas = [(ws, hoja) for ws, hoja in hojas if hoja]
        if not conocidas:
            return _registros(list(wb.active.iter_rows(values_only=True)), "equipos")
        salida = []
        for ws, hoja in conocidas:
            salida += _registros(list(ws.iter_rows(values_only=True)), hoja)
        return salida
    texto = archivo.read().decode("utf-8-sig", errors="replace")
    dialecto = csv.Sniffer().sniff(texto[:2048], delimiters=";,\t") if texto.strip() else csv.excel
    return _registros(list(csv.reader(io.StringIO(texto), dialecto)), "equipos")


def _buscar_tipo(nombre, hoja, errores, fila, crear):
    """El tipo de instrumento debe ser de la familia de la hoja. Con `crear`, el que falte se agrega a esa familia."""
    familia = HOJAS.get(hoja)
    if not nombre:
        if familia:
            errores.append(f"Escribe el tipo de instrumento (en la hoja «{hoja}» es obligatorio).")
        return None
    tipo = TipoInstrumento.objects.filter(nombre__iexact=nombre).first()
    if tipo is None and crear:
        tipo = TipoInstrumento.objects.create(nombre=nombre, familia=familia or Familia.OTRO)
        fila.avisos.append(f"Se agregará el tipo de instrumento «{nombre}» ({tipo.get_familia_display().lower()}).")
    if tipo is None:
        errores.append(f"El tipo de instrumento «{nombre}» no existe en la lista.")
        return None
    if familia and tipo.familia != familia:
        errores.append(f"«{tipo}» es de {tipo.get_familia_display().lower()}: va en otra hoja, no en «{hoja}».")
        return None
    return tipo


# Listas que se pueden completar desde el archivo: campo → (modelo, campo de búsqueda, nombre en el aviso)
AGREGABLES = {"marca": (Marca, "nombre", "la marca"), "ubicacion": (Ubicacion, "nombre", "la ubicación"),
              "unidad": (Unidad, "simbolo", "la unidad")}


def _buscar(modelo, campo, valor, etiqueta, errores, fila=None, crear=False):
    if not valor:
        return ""
    obj = modelo.objects.filter(**{f"{campo}__iexact": valor}).first()
    if obj is None and crear:
        datos = {"nombre": valor, "simbolo": valor} if modelo is Unidad else {"nombre": valor}
        obj = modelo.objects.create(**datos)
        fila.avisos.append(f"Se agregará {etiqueta.lower()} «{valor}» a la lista.")
    if obj is None:
        errores.append(f"{etiqueta} «{valor}» no existe en la lista.")
        return ""
    return obj.pk


def _crear_area(laboratorio, area, usuario, errores, fila):
    """Crea el área y, si hace falta, el laboratorio (solo lo hace el Administrador)."""
    if not laboratorio:
        errores.append(f"El área «{area}» no existe. Escribe también el laboratorio para agregarla.")
        return ""
    lab = Laboratorio.objects.filter(nombre__iexact=laboratorio).first()
    if lab is None:
        lab = Laboratorio.objects.create(nombre=laboratorio)
        fila.avisos.append(f"Se agregará el laboratorio «{laboratorio}».")
    fila.avisos.append(f"Se agregará el área «{area}» a {lab}.")
    return Area.objects.create(laboratorio=lab, nombre=area).pk


def _buscar_area(laboratorio, area, usuario, errores, fila=None, crear=False):
    """El área por su nombre (y el del laboratorio). Si se deja vacía y el usuario tiene una sola área, va a esa."""
    if not area:
        if usuario is not None and not usuario.ve_todas_las_areas and usuario.areas_permitidas().count() == 1:
            return usuario.areas_permitidas().get().pk
        return ""
    qs = Area.objects.filter(nombre__iexact=area)
    if laboratorio:
        qs = qs.filter(laboratorio__nombre__iexact=laboratorio)
    encontradas = list(qs[:2])
    if not encontradas and crear:
        return _crear_area(laboratorio, area, usuario, errores, fila)
    if not encontradas:
        donde = f" del laboratorio «{laboratorio}»" if laboratorio else ""
        errores.append(f"El área «{area}»{donde} no existe en la lista.")
        return ""
    if len(encontradas) > 1:
        errores.append(f"Hay varias áreas «{area}»: escribe también el laboratorio.")
        return ""
    return encontradas[0].pk


def validar(registros, usuario=None, crear=False):
    """Con `crear`, las marcas, ubicaciones, unidades, laboratorios y áreas que falten se crean: llamarla dentro de una
    transacción y deshacerla si solo se está revisando."""
    revisadas = set()
    for reg in registros:
        hoja = reg.get("_hoja", "equipos")
        if hoja in revisadas:
            continue
        revisadas.add(hoja)
        faltan = [c for c in ("codigo", "descripcion", "marca", "modelo", "numero_serie") if c not in reg]
        if faltan:
            return [Fila((hoja, 1), {}, [f"Faltan columnas en la hoja «{hoja}»: {', '.join(faltan)}. Usa la plantilla."])]
    filas, vistos_codigo, vistos_serie = [], {}, {}
    criticidades = {e.lower(): v for v, e in Criticidad.choices}
    for n, reg in enumerate(registros, start=2):
        hoja = reg.get("_hoja", "equipos")
        fila = Fila((hoja, reg.get("_fila", n)), {c: reg.get(c, "") for c in encabezados_de(hoja)})
        e = fila.errores
        d = dict(fila.datos)
        tipo = _buscar_tipo(d.pop("tipo"), hoja, e, fila, crear)
        d["tipo"] = tipo.pk if tipo else ""
        d["critico"] = "on" if d.pop("critico", "").strip().lower() in SI else ""
        for dato in familias.datos(tipo.familia if tipo else None):
            valor, error = familias.limpiar(dato, d.pop(dato.clave, ""))  # acepta «variable» por «Variable»
            if error:
                e.append(error)
            d[f"dt_{dato.clave}"] = valor
        for campo, (modelo, buscar_por, nombre) in AGREGABLES.items():
            d[campo] = _buscar(modelo, buscar_por, d[campo], nombre.capitalize(), e, fila, crear)
        d["area"] = _buscar_area(d.pop("laboratorio"), d["area"], usuario, e, fila, crear)
        d["responsable"] = _buscar(Usuario, "username", d["responsable"], "El usuario responsable", e)
        # Si el nombre no está en la lista, ese error ya lo dice: no se repite «este campo es obligatorio».
        fallidos = {c for c in (*AGREGABLES, "area", "responsable") if fila.datos[c] and not d[c]}
        d["criticidad"] = criticidades.get((d["criticidad"] or "media").lower(), d["criticidad"])
        for num in ("rango_min", "rango_max", "resolucion", "capacidad"):
            d[num] = d[num].replace(",", ".")
        d["intervalo_meses"] = d["intervalo_meses"] or "12"
        form = EquipoForm(data=d, usuario=usuario)
        if not form.is_valid():
            for campo, errs in form.errors.items():
                if campo in fallidos:
                    continue
                etiqueta = form.fields[campo].label if campo in form.fields else ""
                e.extend(f"{etiqueta}: {x}" if etiqueta else x for x in errs)
        codigo = d["codigo"].strip().upper()
        if codigo in vistos_codigo:
            e.append(f"El código {codigo} se repite en {vistos_codigo[codigo]} del archivo.")
        vistos_codigo.setdefault(codigo, fila.lugar)
        clave = (d["marca"], d["modelo"].strip().lower(), d["numero_serie"].strip().lower())
        if all(clave) and clave in vistos_serie:
            e.append(f"Marca, modelo y n.º de serie se repiten en {vistos_serie[clave]} del archivo.")
        vistos_serie.setdefault(clave, fila.lugar)
        fila.datos = d
        filas.append(fila)
    return filas


def guardar(filas, usuario=None):
    creados = []
    for f in filas:
        form = EquipoForm(data=f.datos, usuario=usuario)
        if not form.is_valid():  # Algo cambió desde la revisión: se cancela todo.
            raise ValueError(f"La {f.lugar} dejó de ser válida: {form.errors.as_text()}")
        equipo = form.save(commit=False)
        equipo._change_reason = "Carga masiva desde archivo"
        equipo.save()
        creados.append(equipo)
    return creados


EJEMPLOS = {
    "equipos": ["EQ-0100", "PAT-123", "Termómetro digital", "", "Marca A", "TD-10", "SN-0100"],
    "balanzas": ["EQ-0101", "PAT-124", "Balanza analítica 220 g", "Balanza analítica", "Marca A", "BA-220", "SN-0101"],
    "termicos": ["EQ-0102", "", "Estufa de cultivo 50 L", "Estufa de cultivo", "Marca A", "EC-50", "SN-0102"],
    "cabinas": ["EQ-0103", "", "Cabina de bioseguridad clase II", "Cabina de bioseguridad", "Marca A", "CB-4", "SN-0103"],
    "lectores": ["EQ-0104", "", "Lector de microplaca", "Lector de microplaca", "Marca A", "LM-96", "SN-0104"],
    "pipetas": ["EQ-0105", "", "Pipeta 8 canales 10-100 µL", "Pipeta multicanal", "Marca A", "PM-8", "SN-0105"],
}
EJEMPLO_COMUN = ["Laboratorio Central", "Metrología", "Sala de balanzas", "mperez"]
EJEMPLO_MEDIDA = {
    "equipos": ["°C", "-10", "50", "0,1", "", "±0,5 °C"],
    "balanzas": ["g", "0", "220", "0,0001", "220", ""],
    "termicos": ["°C", "20", "60", "0,1", "", "±1 °C"],
    "cabinas": ["m/s", "0", "1", "0,01", "", ""],
    "lectores": ["DO", "0", "4", "0,001", "", ""],
    "pipetas": ["µL", "10", "100", "0,1", "100", ""],
}
EJEMPLO_DATOS = {
    "balanzas": ["I", "0,0001", "0,001"], "termicos": ["37", "9", "50"], "cabinas": ["II A2", "4 pies"],
    "lectores": ["405; 450; 492; 620", "96"], "pipetas": ["8", "100", "Variable"],
}


def plantilla():
    """Excel de carga: una hoja por familia, con listas desplegables para lo que tiene opciones."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    from apps.core.exportar import libro_excel

    hojas = []
    for hoja in HOJAS:
        fila = EJEMPLOS[hoja] + EJEMPLO_COMUN + EJEMPLO_MEDIDA[hoja] + ["Alta", "Sí" if hoja != "equipos" else "No", "12"]
        hojas.append((hoja, encabezados_de(hoja), [fila + EJEMPLO_DATOS.get(hoja, [])]))
    explicacion = [[c, d] for c, _, d in COLUMNAS]
    for hoja, familia in HOJAS.items():
        for d in familias.datos(familia):
            opciones = f" Opciones: {', '.join(d.opciones)}." if d.opciones else ""
            explicacion.append([d.clave, f"Hoja «{hoja}»: {d.etiqueta}. {d.ayuda}{opciones}".strip()])
    explicacion += [["", ""], ["Hojas", "Usa la hoja de la familia del instrumento; los demás van en «equipos». "
                                        "Borra las filas de ejemplo antes de cargar."]]
    hojas.append(("instrucciones", ["Columna", "Qué escribir"], explicacion))

    # Lista de tipos por familia en una hoja oculta, para los desplegables.
    columnas_tipo = {}
    listas = []
    for i, (hoja, familia) in enumerate(HOJAS.items()):
        qs = TipoInstrumento.objects.filter(activo=True).order_by("nombre")
        nombres = [t.nombre for t in (qs.filter(familia=familia) if familia else qs)]
        columnas_tipo[hoja] = (get_column_letter(i + 1), max(len(nombres), 1))
        listas.append(nombres)
    alto = max(len(n) for n in listas) if listas else 0
    hojas.append(("listas", list(HOJAS), [[(n[j] if j < len(n) else "") for n in listas] for j in range(alto)]))
    wb = libro_excel(hojas)
    wb["listas"].sheet_state = "hidden"

    for hoja in HOJAS:
        ws = wb[hoja]
        cols = encabezados_de(hoja)

        def lista(columna, formula):
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            letra = get_column_letter(cols.index(columna) + 1)
            dv.add(f"{letra}2:{letra}1000")

        letra, n = columnas_tipo[hoja]
        lista("tipo", f"=listas!${letra}$2:${letra}${n + 1}")
        lista("criticidad", '"Alta,Media,Baja"')
        lista("critico", '"Sí,No"')
        for d in familias.datos(HOJAS[hoja]):
            if d.opciones:
                lista(d.clave, '"' + ",".join(d.opciones) + '"')
    return wb
